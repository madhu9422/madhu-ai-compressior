from io import BytesIO
from docx import Document
from docx.shared import Inches

def create_docx_report(name_a: str, name_b: str, results: dict) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    doc.add_heading("AI Document Comparison Report", 0)
    doc.add_paragraph("Informational comparison generated from extracted PDF text.")
    doc.add_heading("Document A", level=1)
    doc.add_paragraph(f"File: {name_a}")
    doc.add_paragraph(results.get("summary_a", "Not found in the provided text"))
    doc.add_heading("Document B", level=1)
    doc.add_paragraph(f"File: {name_b}")
    doc.add_paragraph(results.get("summary_b", "Not found in the provided text"))
    doc.add_heading("Similarities", level=1)
    _add_list(doc, results.get("similarities", []))
    doc.add_heading("Differences", level=1)
    _add_list(doc, results.get("differences", []))
    doc.add_heading("Key Details Comparison", level=1)
    details = results.get("key_details", [])
    if details:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Shading Accent 1"
        for cell, heading in zip(table.rows[0].cells, ["Field", "Document A", "Document B", "Change"]):
            cell.text = heading
        for item in details:
            cells = table.add_row().cells
            for cell, key in zip(cells, ["field", "document_a", "document_b", "change"]):
                cell.text = str(item.get(key, "Not found in the provided text"))
    else:
        doc.add_paragraph("No key details returned.")
    doc.add_heading("Overall Conclusion", level=1)
    doc.add_paragraph(results.get("overall_conclusion", "Not found in the provided text"))
    doc.add_heading("Limitations", level=1)
    _add_list(doc, results.get("limitations", []))
    doc.add_paragraph("Note: This report is informational. It is not legal advice and does not establish compliance or equivalence.")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()

def _add_list(doc, items):
    if not items:
        doc.add_paragraph("None identified.")
        return
    for item in items:
        doc.add_paragraph(str(item), style="List Bullet")
